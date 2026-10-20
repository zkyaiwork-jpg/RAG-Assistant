"""Generate the extended policy corpus as DOCX files.

Content lives in ``corpus_batch_*.py``; this module only handles layout and numbering
so every clause gets a stable 第N条 prefix without hand-maintained ordinals.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from corpus_batch_1 import DOCS as BATCH_1
from corpus_batch_2 import DOCS as BATCH_2
from corpus_batch_3 import DOCS as BATCH_3
from corpus_batch_4 import DOCS as BATCH_4


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "samples" / "corpus"
META = "发布部门：公司行政部　　生效日期：2026 年 10 月 1 日　　版本：V1.0"
BODY_FONT = "宋体"
HEAD_FONT = "黑体"
DIGITS = "零一二三四五六七八九"

ALL_DOCS = [*BATCH_1, *BATCH_2, *BATCH_3, *BATCH_4]


def chinese_number(value: int) -> str:
    if value < 10:
        return DIGITS[value]
    if value == 10:
        return "十"
    if value < 20:
        return "十" + DIGITS[value % 10]
    tens, ones = divmod(value, 10)
    return DIGITS[tens] + "十" + (DIGITS[ones] if ones else "")


def set_run_font(run, name: str, size: int, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def build_document(entry: dict) -> tuple[int, int]:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    for style_name, size in (("Heading 1", 14),):
        style = document.styles[style_name]
        style.font.name = HEAD_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x2A, 0x37)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)

    section = document.sections[0]
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run(entry["title"]), HEAD_FONT, 20, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run(entry["subtitle"]), BODY_FONT, 11)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run(META), BODY_FONT, 10)

    number = 0
    characters = 0
    for heading, clauses in entry["sections"]:
        document.add_heading(heading, level=1)
        for clause in clauses:
            number += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Pt(24)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.5
            text = f"第{chinese_number(number)}条　{clause}"
            characters += len(text)
            set_run_font(paragraph.add_run(text), BODY_FONT, 12)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / entry["filename"]
    document.save(path)
    return characters, number


def main() -> int:
    total_characters = 0
    total_clauses = 0
    print(f"{'文档':<28}{'字符':>6}{'条款':>6}  文件夹")
    print("-" * 60)
    for entry in ALL_DOCS:
        characters, clauses = build_document(entry)
        total_characters += characters
        total_clauses += clauses
        print(f"{entry['filename']:<28}{characters:>6}{clauses:>6}  {entry['folder']}")
    print("-" * 60)
    print(f"共 {len(ALL_DOCS)} 篇，{total_clauses} 条，{total_characters} 字符，平均 {total_characters // len(ALL_DOCS)} 字符/篇")
    print(f"输出目录：{OUTPUT_DIR}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
